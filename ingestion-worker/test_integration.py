"""Explicit opt-in fixtures: localhost:13316/knowledge_ingest_test, never the business DB."""
import io
import json
import os
import sys
import time
import uuid
import subprocess
import logging
from pathlib import Path

import pytest
import pymysql
from minio import Minio

import ingest_runtime as runtime
import worker

pytestmark = pytest.mark.skipif(os.getenv('INGEST_INTEGRATION') != 'true', reason='requires isolated integration stack')


@pytest.fixture(autouse=True)
def fixtures(monkeypatch):
    monkeypatch.setattr(worker, 'MYSQL', dict(host='127.0.0.1', port=13316, user='root', password='integration',
        database='knowledge_ingest_test', autocommit=True, charset='utf8mb4', cursorclass=pymysql.cursors.DictCursor))
    monkeypatch.setattr(worker, 'MINIO', Minio('127.0.0.1:19000',access_key='minioadmin',secret_key='minioadmin',secure=False))
    monkeypatch.setattr(worker, 'BUCKET', 'ingestion-test')
    monkeypatch.setattr(worker, 'QDRANT', 'http://127.0.0.1:16333')
    monkeypatch.setattr(worker, 'COLLECTION', 'knowledge_documents_test')
    monkeypatch.setattr(worker, 'DIMENSIONS', 4)
    monkeypatch.setattr(worker, 'MINERU_URL', '')
    monkeypatch.setattr(worker, 'MODEL_URL', '')
    monkeypatch.setattr(worker, 'VISION_URL', '')
    # Model boundary is a deterministic stub; MySQL/MinIO/Qdrant/RocketMQ are real.
    monkeypatch.setattr(worker, 'hybrid_embed', lambda texts: [{'dense':[1.,0.,0.,0.], 'sparse':{'indices':[], 'values':[]}} for _ in texts])
    runtime.STOP.clear()
    if not worker.MINIO.bucket_exists(worker.BUCKET): worker.MINIO.make_bucket(worker.BUCKET)


def sql(query, args=()):
    with worker.db() as connection:
        with connection.cursor() as cursor:
            cursor.execute(query,args)
            return cursor.fetchall()


def create_job(data=b'# Heading\n\nTest text.', extension='md'):
    key=str(uuid.uuid4())+'.'+extension
    worker.MINIO.put_object(worker.BUCKET,key,io.BytesIO(data),len(data))
    with worker.db() as connection:
        with connection.cursor() as cursor:
            cursor.execute("INSERT INTO knowledge_document(title,original_filename,content_type,file_size,object_key,created_by) VALUES(%s,%s,'application/octet-stream',%s,%s,1)", (key,key,len(data),key))
            document_id=cursor.lastrowid
            cursor.execute('INSERT INTO knowledge_ingest_job(document_id) VALUES(%s)',(document_id,))
            job_id=cursor.lastrowid
    return dict(schemaVersion=1,eventId=str(uuid.uuid4()),jobId=job_id,documentId=document_id,generation=1)


@pytest.mark.parametrize('extension', ['md','docx','pdf'])
def test_formats_and_replay(extension,mq):
    if extension=='docx':
        from docx import Document
        stream=io.BytesIO(); doc=Document(); doc.add_paragraph('Document test'); doc.save(stream); data=stream.getvalue()
    elif extension=='pdf':
        from reportlab.pdfgen.canvas import Canvas
        stream=io.BytesIO(); canvas=Canvas(stream); canvas.drawString(60,700,'PDF test'); canvas.save(); data=stream.getvalue()
    else: data=b'# Heading\n\nDocument test.'
    event=create_job(data,extension)
    consumer,producer,_=mq
    send(producer,event)
    runtime.handle(consumer,receive_event(consumer,event),worker)
    job=sql('SELECT * FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]
    assert job['status']==('DEGRADED' if extension=='pdf' else 'SUCCESS')
    assert sql('SELECT * FROM knowledge_chunk WHERE document_id=%s',(event['documentId'],))
    assert runtime.process(event,worker)
    assert sql('SELECT execution_attempts FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]['execution_attempts']==1


def test_three_failures_and_manual_generation(monkeypatch):
    event=create_job()
    def fail(_): raise RuntimeError('temporary model failure')
    original=worker.hybrid_embed
    monkeypatch.setattr(worker,'hybrid_embed',fail)
    assert not runtime.process(event,worker)
    assert not runtime.process(event,worker)
    assert runtime.process(event,worker)
    assert sql('SELECT status FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]['status']=='FAILED'
    sql("UPDATE knowledge_ingest_job SET generation=2,status='PENDING',execution_attempts=0 WHERE id=%s",(event['jobId'],))
    assert runtime.process(event,worker)  # Old delivery must not execute new generation.
    event['generation']=2
    monkeypatch.setattr(worker,'hybrid_embed',original)
    assert runtime.process(event,worker)


def test_migration_idempotent():
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
    from migrate_knowledge_queue import backfill
    event=create_job()
    with worker.db() as connection:
        backfill(connection,True); backfill(connection,True)
    assert sql('SELECT COUNT(*) AS n FROM knowledge_ingest_outbox WHERE job_id=%s',(event['jobId'],))[0]['n']==1


def retry_limit(limit):
    subprocess.run(['docker','exec','knowledge-mq-test-rocketmq-broker-1','sh','mqadmin',
        'updateSubGroup','-n','rocketmq-namesrv:9876','-c','KnowledgeCluster',
        '-g','knowledge-ingest-worker','-r',str(limit),'-m','true'],check=True,capture_output=True,timeout=30)


@pytest.fixture
def mq(request):
    from rocketmq import ClientConfiguration, Credentials, FilterExpression, SimpleConsumer, Producer
    is_dead_test=request.node.name=='test_real_dead_letter_updates_failed_status'
    if is_dead_test: retry_limit(1)
    config=ClientConfiguration('localhost:18081',Credentials(),request_timeout=10)
    consumer=SimpleConsumer(config,'knowledge-ingest-worker',{'knowledge-ingest':FilterExpression()},await_duration=1)
    producer=Producer(config,{'knowledge-ingest'})
    dead=SimpleConsumer(config,'knowledge-ingest-worker-dead',{'%DLQ%knowledge-ingest-worker':FilterExpression()},await_duration=1)
    try:
        consumer.startup(); producer.startup(); dead.startup()
        yield consumer,producer,dead
    finally:
        consumer.shutdown(); producer.shutdown(); dead.shutdown()
        if is_dead_test: retry_limit(16)


def receive_recovering(consumer):
    from rocketmq.v5.exception.client_exception import ClientException
    try:
        return consumer.receive(1,10) or []
    except ClientException as error:
        # Same recovery as runtime.main: a restarted Broker can temporarily lack routes.
        logging.warning('Transient receive failure during fault test: %s',error)
        time.sleep(.2)
        return []


def ack_recovering(consumer,message):
    from rocketmq.v5.exception.client_exception import BadRequestException
    try:
        consumer.ack(message)
        return True
    except BadRequestException as error:
        if error.code != 40013: raise
        # An expired receipt must be replaced by redelivery; never count it as an ACK.
        logging.warning('Expired receipt during fault test: delivery=%s',message.delivery_attempt)
        return False


def receive_event(consumer,event,seconds=50):
    deadline=time.monotonic()+seconds
    while time.monotonic()<deadline:
        for message in consumer.receive(1,10) or []:
            if json.loads(message.body).get('eventId')==event['eventId']: return message
            consumer.ack(message)  # Only isolated test topic contains fixtures.
    raise AssertionError('message not delivered')


def send(producer,event):
    from rocketmq import Message
    message=Message(); message.topic='knowledge-ingest'; message.body=json.dumps(event).encode()
    producer.send(message)


def test_real_mq_ack_loss_redelivery_and_renewal(mq):
    consumer,producer,dead=mq
    event=create_job()
    send(producer,event)
    message=receive_event(consumer,event)
    consumer.change_invisible_duration(message,10)
    assert runtime.process(event,worker)
    # Simulate worker death after commit, before ACK: same event must return without re-execution.
    redelivery=receive_event(consumer,event)
    runtime.handle(consumer,redelivery,worker)
    assert sql('SELECT execution_attempts FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]['execution_attempts']==1
    # The DLQ consumer has been started against its actual broker route by the fixture.


def test_worker_process_death_releases_lock_and_redelivers(mq, tmp_path):
    consumer,producer,_=mq
    event=create_job()
    send(producer,event)
    marker=tmp_path/'executing'
    child_code='''
import sys,time,json
from pathlib import Path
import worker,ingest_runtime as runtime
from rocketmq import ClientConfiguration,Credentials,SimpleConsumer,FilterExpression
consumer=SimpleConsumer(ClientConfiguration('localhost:18081',Credentials(),request_timeout=10),'knowledge-ingest-worker',{'knowledge-ingest':FilterExpression()},await_duration=1)
consumer.startup()
def pause(state):
    Path(sys.argv[2]).write_text('locked',encoding='utf-8')
    time.sleep(120)
worker.pipeline=type('Pipeline',(),{'invoke':staticmethod(pause)})()
deadline=time.monotonic()+60
while time.monotonic()<deadline:
    for message in consumer.receive(1,10) or []:
        if json.loads(message.body)['eventId']==sys.argv[1]:
            runtime.handle(consumer,message,worker)
        else: consumer.ack(message)
'''
    env=dict(os.environ, MYSQL_HOST='127.0.0.1', MYSQL_PORT='13316', MYSQL_USER='root',
             MYSQL_PASSWORD='integration', MYSQL_DATABASE='knowledge_ingest_test')
    child=subprocess.Popen([sys.executable,'-c',child_code,event['eventId'],str(marker)],
                           cwd=Path(__file__).parent,env=env)
    try:
        deadline=time.monotonic()+60
        while not marker.exists() and time.monotonic()<deadline:
            assert child.poll() is None
            time.sleep(.2)
        assert marker.exists(), 'child did not acquire task'
    finally:
        child.kill(); child.wait(timeout=10)
    message=receive_event(consumer,event,60)
    runtime.handle(consumer,message,worker)
    job=sql('SELECT status,execution_attempts FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]
    assert job=={'status':'SUCCESS','execution_attempts':2}


def test_broker_restart_preserves_unacknowledged_message(mq):
    consumer,producer,_=mq
    event=create_job()
    send(producer,event)
    receive_event(consumer,event)
    subprocess.run(['docker','restart','knowledge-mq-test-rocketmq-broker-1'],check=True,capture_output=True,timeout=60)
    deadline=time.monotonic()+90
    message=None
    while time.monotonic()<deadline and message is None:
        try: message=receive_event(consumer,event,15)
        except Exception: time.sleep(1)
    assert message is not None
    runtime.handle(consumer,message,worker)
    assert sql('SELECT status FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]['status']=='SUCCESS'


def test_image_replay_and_shorter_index_removes_old_vectors(monkeypatch):
    from docx import Document
    from PIL import Image
    import requests
    stream=io.BytesIO(); image=io.BytesIO(); Image.new('RGB',(10,10),'red').save(image,format='PNG'); image.seek(0)
    doc=Document(); doc.add_paragraph('Some image context. '*20); doc.add_picture(image); doc.save(stream)
    monkeypatch.setattr(worker,'MAX_CHUNK_CHARS',60)
    event=create_job(stream.getvalue(),'docx')
    assert runtime.process(event,worker)
    assert sql('SELECT COUNT(*) AS n FROM knowledge_chunk WHERE document_id=%s',(event['documentId'],))[0]['n']>1
    # Same image record can be written again without violating the document/hash unique key.
    for generation in (2,3):
        sql("UPDATE knowledge_ingest_job SET generation=%s,status='PENDING',execution_attempts=0 WHERE id=%s",(generation,event['jobId']))
        event['generation']=generation
        monkeypatch.setattr(worker,'MAX_CHUNK_CHARS',10000)
        assert runtime.process(event,worker)
    assert sql('SELECT COUNT(*) AS n FROM knowledge_document_image WHERE document_id=%s',(event['documentId'],))[0]['n']==1
    result=requests.post(worker.QDRANT+'/collections/'+worker.COLLECTION+'/points/count',json={
        'filter':{'must':[{'key':'documentId','match':{'value':event['documentId']}}]},'exact':True},timeout=10)
    result.raise_for_status()
    assert result.json()['result']['count']==1


def test_real_dead_letter_updates_failed_status(mq):
    consumer,producer,dead=mq
    event=create_job()
    send(producer,event)
    # Continue polling the original subscription: Proxy evaluates exhausted deliveries
    # while receiving from the retry queue, rather than on a passive timeout alone.
    deadline=time.monotonic()+90
    message=None
    deliveries=[]
    while time.monotonic()<deadline and message is None:
        for item in receive_recovering(consumer):
            if json.loads(item.body).get('eventId')==event['eventId']:
                deliveries.append(item.delivery_attempt)
            else: ack_recovering(consumer,item)
        for item in receive_recovering(dead):
            if json.loads(item.body).get('eventId')==event['eventId']: message=item
            else: ack_recovering(dead,item)
    assert message is not None, f'no dead letter after deliveries {deliveries}'
    runtime.handle(dead,message,worker,dead_letter=True)
    job=sql('SELECT status,execution_attempts FROM knowledge_ingest_job WHERE id=%s',(event['jobId'],))[0]
    assert job=={'status':'FAILED','execution_attempts':0}


def test_first_consumer_receives_entire_backlog():
    from rocketmq import ClientConfiguration, Credentials, FilterExpression, SimpleConsumer, Producer, Message
    suffix=uuid.uuid4().hex[:10]
    topic='knowledge-bootstrap-'+suffix
    group=topic+'-worker'
    broker='knowledge-mq-test-rocketmq-broker-1'
    script=str(Path(__file__).resolve().parents[1]/'docker'/'rocketmq'/'init.sh')
    subprocess.run(['docker','cp',script,broker+':/tmp/knowledge-init.sh'],check=True,capture_output=True)
    command=['docker','exec','-e','ROCKETMQ_TOPIC='+topic,'-e','ROCKETMQ_CONSUMER_GROUP='+group,
             '-e','ROCKETMQ_INIT_STATE_DIR=/tmp/knowledge-bootstrap-test',broker,'sh','/tmp/knowledge-init.sh']
    initial=subprocess.run(command,check=True,capture_output=True,text=True,timeout=60)
    assert '#offset' in initial.stdout
    config=ClientConfiguration('localhost:18081',Credentials(),request_timeout=10)
    producer=Producer(config,{topic})
    consumer=SimpleConsumer(config,group,{topic:FilterExpression()},await_duration=1)
    producer.startup()
    try:
        for number in range(20):
            message=Message(); message.topic=topic; message.body=str(number).encode()
            producer.send(message)
        # No consumer existed while these messages were published.
        consumer.startup()
        received=set()
        deadline=time.monotonic()+90
        while len(received)<20 and time.monotonic()<deadline:
            for message in receive_recovering(consumer):
                if ack_recovering(consumer,message): received.add(int(message.body))
        assert received==set(range(20))
        repeated=subprocess.run(command,check=True,capture_output=True,text=True,timeout=60)
        assert '#offset' not in repeated.stdout
        assert repeated.stdout.count('Offsets already initialized:')==2
    finally:
        producer.shutdown(); consumer.shutdown()
